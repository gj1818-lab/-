"""주간보고 아티팩트('27SS 소재 진척현황') 데이터 생성.

사용: python3 build_weekly_report_data.py <운영시트.xlsx> <EDW.xlsx> <기준일 YYYY-MM-DD> <출력.json>
운영시트 = '27SS_생진테+소재DB' xlsx export, EDW = '27SS 생산 프로세스 진척현황 (자동화)' xlsx export.
문제 스타일 = 소재역산 트래커 V열(생산처-원단처 협의날짜) 빨간 칠(FFFF0000).
"""
import datetime
import json
import re
import sys
import warnings

from openpyxl import load_workbook

warnings.filterwarnings('ignore')

oper_path, edw_path, base_str, out_path = sys.argv[1:5]
base = datetime.date.fromisoformat(base_str)

# 열 번호 (소재역산 트래커, 4행 헤더) — 팀이 열을 삽입하면 밀리므로 헤더로 검증
COL = dict(code=5, name=6, season=8, form=9, ship=12, deadline=15, bisugi=17,
           jungbal=20, hyup_date=22, y_ship=25, fabric=26, matin=27, labdip=28,
           factory=30, agency=33)
EXPECT = {5: 'Style', 8: '시즌', 15: '소재입고', 17: '비수기', 20: '정발주',
          22: '생산처-원단처', 25: '소재', 28: 'LAB DIP', 30: '생산처'}

wb = load_workbook(oper_path, data_only=True)
ws = wb['소재역산 트래커']
for c, k in EXPECT.items():
    h = str(ws.cell(4, c).value or '')
    assert h.startswith(k), f'열 구조 변경: {c}열 헤더 {h!r} (기대 {k!r})'

wbE = load_workbook(edw_path, data_only=True)
wsE = wbE['EDW_RAW']
h0 = [str(c.value) for c in next(wsE.iter_rows(min_row=1, max_row=1))][:11]
assert h0[0] == 'StyleCode' and h0[7] == '생산MD' and h0[8] == '기획MD' and h0[10] == '디자인담당', h0

tot, il, ppl = {}, {}, {}
for row in wsE.iter_rows(min_row=2):
    a = row[0].value
    if not a:
        continue
    a = str(a).strip()
    if not a or a == 'StyleCode':
        continue
    tot[a] = tot.get(a, 0) + 1
    if row[17].value and str(row[17].value).strip():
        il[a] = il.get(a, 0) + 1
    e = ppl.setdefault(a, {'prod': set(), 'plan': set(), 'design': set()})
    for k, idx in (('prod', 7), ('plan', 8), ('design', 10)):
        v = row[idx].value
        if v and str(v).strip():
            e[k].add(str(v).strip())


def jak(c):
    tt = tot.get(c, 0)
    if tt == 0:
        return '－'
    x = il.get(c, 0)
    return 'O' if x == tt else ('부분' if x > 0 else 'X')


def fdate(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return None


def fstr(v):
    d = fdate(v)
    if d:
        return d.strftime('%m/%d')
    return str(v).strip() if v is not None else ''


def cell(r, k):
    return ws.cell(row=r, column=COL[k]).value


data, red = [], set()
for r in range(5, ws.max_row + 1):
    code = cell(r, 'code')
    if not code:
        continue
    code = str(code).strip()
    f = ws.cell(row=r, column=COL['hyup_date']).fill
    if f and f.patternType == 'solid' and getattr(f.start_color, 'rgb', None) == 'FFFF0000':
        red.add(code)
    spring = str(cell(r, 'season') or '').startswith('S')
    bis = str(cell(r, 'bisugi') or '').strip() == 'O'
    if not (spring or bis):
        continue
    dl = fdate(cell(r, 'deadline'))
    ydt = fdate(cell(r, 'y_ship'))
    if ydt is None and isinstance(cell(r, 'y_ship'), str):
        m = re.findall(r'(\d{1,2})\s*[/.]\s*(\d{1,2})', cell(r, 'y_ship'))
        if m:
            mo, dy = int(m[-1][0]), int(m[-1][1])
            ydt = datetime.date(2026 if mo >= 9 else 2027, mo, dy)
    arr = ydt + datetime.timedelta(days=15) if ydt else None
    jb = str(cell(r, 'jungbal') or '').strip()
    p = ppl.get(code, {})
    data.append({
        'code': code, 'name': fstr(cell(r, 'name')),
        'spring': spring, 'bisugi': bis, 'prob': code in red,
        'plan': ' / '.join(sorted(p.get('plan', []))) or '미지정',
        'design': ' / '.join(sorted(p.get('design', []))) or '미지정',
        'prod': ' / '.join(sorted(p.get('prod', []))) or '미지정',
        'form': fstr(cell(r, 'form')),
        'ship': fstr(cell(r, 'ship')),
        'deadline': dl.strftime('%m/%d') if dl else '',
        'dday': (dl - base).days if dl else None,
        'jakji': jak(code),
        'jungbal': 'O' if jb in ('O', 'ㅇ') else ('예정' if jb == '발주예정' else '미확정'),
        'labdip': fstr(cell(r, 'labdip')),
        'factory': fstr(cell(r, 'factory')),
        'agency': fstr(cell(r, 'agency')),
        'arrive': arr.strftime('%m/%d') if arr else '',
        'arrive_late': bool(arr and dl and arr > dl),
        'fabric': fstr(cell(r, 'fabric')),
        'matin': fstr(cell(r, 'matin')),
        'hyupui': fstr(cell(r, 'y_ship'))})

summer_excl = sorted(red - {x['code'] for x in data})
data.sort(key=lambda x: (x['code'].startswith('8'), x['deadline'] or '99/99', x['code']))

prob = [x for x in data if x['prob']]
rest = [x for x in prob if not x['code'].startswith('8')]
blue = [x for x in prob if x['code'].startswith('8')]
s_all = [x for x in data if x['spring']]
b_all = [x for x in data if x['bisugi']]
s_gen = [x for x in s_all if not x['code'].startswith('8')]


def agg(rows):
    blank = lambda k: sum(1 for x in rows if not x[k].strip())
    return dict(
        urgent=sum(1 for x in rows if x['deadline'] and x['deadline'] <= '10/10'),
        mid=sum(1 for x in rows if x['deadline'] and '10/10' < x['deadline'] <= '11/09'),
        late=sum(1 for x in rows if x['deadline'] and x['deadline'] > '11/09'),
        overdue=sum(1 for x in rows if x['dday'] is not None and x['dday'] < 0),
        jkx=[x['code'] for x in rows if x['jakji'] == 'X'],
        jb=sum(1 for x in rows if x['jungbal'] == '미확정'),
        jbu=sum(1 for x in rows if x['jungbal'] == '미확정' and x['deadline'] and x['deadline'] <= '10/10'),
        ld=blank('labdip'), fab=blank('fabric'), hyu=blank('hyupui'), mat=blank('matin'),
        arrh=sum(1 for x in rows if x['arrive']),
        arrl=sum(1 for x in rows if x['arrive_late']),
        mind=min((x['deadline'] for x in rows if x['deadline']), default=''),
        mindd=min((x['dday'] for x in rows if x['dday'] is not None), default=None))


n = {'prob': len(prob), 'rest': len(rest), 'blue': len(blue),
     's_prob': sum(x['prob'] for x in s_all), 's_all': len(s_all),
     's_gen_prob': sum(x['prob'] for x in s_gen), 's_gen': len(s_gen),
     'b_prob': sum(x['prob'] for x in b_all), 'b_all': len(b_all)}
out = {'base': base_str, 'data': data, 'summer_excl': summer_excl, 'n': n,
       'rest_agg': agg(rest), 'blue_agg': agg(blue)}
json.dump(out, open(out_path, 'w'), ensure_ascii=False)
print('n:', n)
print('일반:', out['rest_agg'])
print('피싱:', out['blue_agg'])
print('여름 제외:', summer_excl)
