#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""27SS 판정 백데이터 시트 (구글시트 변환용 xlsx)
운영 시트가 IMPORTRANGE 한 줄로 끌어갈 수 있는 판정값(작지 일매/사양확정/LAB DIP CFM)을
실시간 수식으로 계산해 두는 소형 시트. 셀 쓰기 커넥터 부재 우회용.
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, 'data', 'source_0808_style_master.csv')
OUT = os.path.join(BASE, 'output', '27SS_판정_백데이터.xlsx')

EDW_URL = "https://docs.google.com/spreadsheets/d/1iSDzZFUv3855k9Zh-tNZALH8z3gY7dTw9j8RlHpJETg"
SOJE_URL = "https://docs.google.com/spreadsheets/d/1_LNPz9XGRmRAJ9Y4cqtFF8NtZTLSIld4vtzLJlOcka0"

with open(SRC, newline='', encoding='utf-8') as f:
    rows = list(csv.DictReader(f))
codes = [r['StyleCode'].strip() for r in rows if r['StyleCode'].strip()]
assert len(codes) == len(set(codes)), "품번 중복"
print("styles:", len(codes))

wb = Workbook()
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="2C5F2D")
thin = Border(*[Side(style='thin', color='CCCCCC')]*4)

# ---- 안내 탭 ----
ws = wb.active; ws.title = "안내"
ws.column_dimensions['A'].width = 110
lines = [
 "27SS 판정 백데이터 시트 — 운영 시트(27SS_생진테+소재DB)가 끌어다 쓰는 계산 전용 시트입니다.",
 "",
 "■ 이 시트가 하는 일",
 "  · 작지발행(일매·디자인담당 결재, EDW R열) / 사양확정(EDW U열) 판정을 SUMPRODUCT+LEN 방식으로 계산",
 "    (IMPORTRANGE 빈 셀을 O로 오집계하던 COUNTIFS 버그 수정판)",
 "  · LAB DIP CFM: 소재DB Merged Data AK열, Drop 제외 최근 컨펌일",
 "  · EDW·소재DB가 갱신되면 이 시트도, 이 시트를 참조하는 운영 시트도 자동 갱신됩니다.",
 "",
 "■ 최초 1회 설정 (액세스 허용 2번)",
 "  1) EDW_RAW 탭 A2 셀 선택 → '액세스 허용' 클릭",
 "  2) 소재DB_RAW 탭 A3 셀 선택 → '액세스 허용' 클릭",
 "",
 "■ 운영 시트 연결 (Q5·AA5에 수식 1개씩 — 별도 안내 참조)",
 "  · 운영 시트에서 최초 1회 이 시트에 대한 '액세스 허용' 필요",
 "",
 "■ 주의: 판정 탭의 수식·품번은 수정하지 마세요. 품번 목록은 0808 원본 424 스타일 기준입니다.",
]
for i, t in enumerate(lines, 1):
    ws[f'A{i}'] = t
ws['A1'].font = Font(bold=True, size=12, color="2C5F2D")

# ---- 판정 탭 ----
ws = wb.create_sheet("판정")
hdr = ["품번", "작지발행(일매)", "사양확정", "LAB DIP CFM", "LAB DIP(텍스트)"]
for c, t in enumerate(hdr, 1):
    cell = ws.cell(row=1, column=c, value=t); cell.font = H; cell.fill = HF
    cell.alignment = Alignment(horizontal='center')
ws.column_dimensions['A'].width = 14
for col in 'BCDE': ws.column_dimensions[col].width = 14
last = len(codes) + 1  # 데이터 마지막 행
for i, code in enumerate(codes):
    rw = i + 2
    ws.cell(row=rw, column=1, value=code)
    ws[f'B{rw}'] = (f'=IF(COUNTIF(EDW_RAW!$A$2:$A$1500,$A{rw})=0,"",'
        f'IF(SUMPRODUCT((EDW_RAW!$A$2:$A$1500=$A{rw})*(LEN(EDW_RAW!$R$2:$R$1500)>0))'
        f'=COUNTIF(EDW_RAW!$A$2:$A$1500,$A{rw}),"O",'
        f'IF(SUMPRODUCT((EDW_RAW!$A$2:$A$1500=$A{rw})*(LEN(EDW_RAW!$R$2:$R$1500)>0))>0,"부분","X")))')
    ws[f'C{rw}'] = (f'=IF(COUNTIF(EDW_RAW!$A$2:$A$1500,$A{rw})=0,"",'
        f'IF(SUMPRODUCT((EDW_RAW!$A$2:$A$1500=$A{rw})*(LEN(EDW_RAW!$U$2:$U$1500)>0))'
        f'=COUNTIF(EDW_RAW!$A$2:$A$1500,$A{rw}),"O",'
        f'IF(SUMPRODUCT((EDW_RAW!$A$2:$A$1500=$A{rw})*(LEN(EDW_RAW!$U$2:$U$1500)>0))>0,"부분","X")))')
    ws[f'D{rw}'] = (f'=IFERROR(MAX(FILTER(소재DB_RAW!$AK$3:$AK$2000,'
        f'(소재DB_RAW!$G$3:$G$2000=$A{rw})*ISERROR(SEARCH("drop",소재DB_RAW!$C$3:$C$2000))'
        f'*(소재DB_RAW!$AK$3:$AK$2000<>""))),"")')
    ws[f'D{rw}'].number_format = 'yyyy-mm-dd'
    ws[f'E{rw}'] = f'=IF($D{rw}="","",TEXT($D{rw},"yyyy-mm-dd"))'
    for col in range(1, 6):
        ws.cell(row=rw, column=col).border = thin

# 검증 카운터 (G:H)
ws['G1'] = "검증 카운터 (자동)"; ws['G1'].font = Font(bold=True, color="2C5F2D")
checks = [
    ("작지 O",   f'=COUNTIF($B$2:$B${last},"O")'),
    ("작지 부분", f'=COUNTIF($B$2:$B${last},"부분")'),
    ("작지 X",   f'=COUNTIF($B$2:$B${last},"X")'),
    ("작지 공란(EDW없음)", f'=SUMPRODUCT(--(LEN($B$2:$B${last})=0))'),
    ("사양 O",   f'=COUNTIF($C$2:$C${last},"O")'),
    ("사양 부분", f'=COUNTIF($C$2:$C${last},"부분")'),
    ("사양 X",   f'=COUNTIF($C$2:$C${last},"X")'),
    ("사양 공란(EDW없음)", f'=SUMPRODUCT(--(LEN($C$2:$C${last})=0))'),
    ("LAB DIP 보유", f'=COUNTIF($E$2:$E${last},"?*")'),
    ("품번 수", f'=COUNTA($A$2:$A${last})'),
]
ws.column_dimensions['G'].width = 20
for i, (label, formula) in enumerate(checks, 2):
    ws[f'G{i}'] = label; ws[f'H{i}'] = formula
ws.freeze_panes = 'A2'

# ---- EDW_RAW 탭 ----
ws = wb.create_sheet("EDW_RAW")
ws['A1'] = "27SS 생산 프로세스 진척현황(자동화) > EDW_RAW 실시간 연동. 최초 1회 A2 셀에서 '액세스 허용' 클릭. 이 탭 수정 금지."
ws['A1'].font = Font(bold=True, color="CC0000")
ws['A2'] = f'=IMPORTRANGE("{EDW_URL}","EDW_RAW!A2:U1500")'

# ---- 소재DB_RAW 탭 ----
ws = wb.create_sheet("소재DB_RAW")
ws['A1'] = "27SS_소재DB > Merged Data 실시간 연동. 최초 1회 A3 셀에서 '액세스 허용' 클릭. 이 탭 수정 금지."
ws['A1'].font = Font(bold=True, color="CC0000")
ws['A3'] = f'=IMPORTRANGE("{SOJE_URL}","Merged Data!A3:AN2000")'

wb.save(OUT)
print("saved:", OUT, os.path.getsize(OUT), "bytes")
