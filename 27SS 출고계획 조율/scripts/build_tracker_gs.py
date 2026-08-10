# 27SS 소재일정 역산 트래커 빌드 스크립트
# 원본: data/source_0808_style_master.csv (0808 첨부 파일 집계본)
# 정렬: ① S코드(봄) → ② 비수기 선정(봄 제외) → ③ 나머지(M→X), 각 그룹 내 출고일 주차 순
import csv, re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

BRAND="2C5F2D"; PAPER="F5F5F5"; RED="E8474C"; ORANGE="F18F01"; TOTALG="E8F5E9"; YELLOW="FFFF00"; NAVY="1B3A4B"
F="맑은 고딕"
thin=Side(style='thin', color='D9D9D9'); border=Border(left=thin,right=thin,top=thin,bottom=thin)

def h(ws,cell,text,size=11,bold=True,color="FFFFFF",fill=BRAND):
    ws[cell]=text; ws[cell].font=Font(name=F,size=size,bold=bold,color=color)
    if fill: ws[cell].fill=PatternFill('solid',fgColor=fill)

rows=list(csv.DictReader(open('data/source_0808_style_master.csv')))

def week_key(w):
    m=re.match(r'(\d+)월(\d+)주',w or '')
    return ((int(m.group(1))-9)%12,int(m.group(2))) if m else (99,99)

def sort_key(r):
    # 그룹: 0=S코드(봄, 납기 최우선) / 1=비수기 선정(봄 제외) / 2=나머지(M→X)
    if r['시즌코드']=='S': grp=0
    elif r['비수기']=='O': grp=1
    else: grp=2
    season_order={'M':0,'X':1}.get(r['시즌코드'],2) if grp==2 else 0
    return (grp, season_order, week_key(r['출고일(주차)']), r['StyleCode'])
rows.sort(key=sort_key)

wb=Workbook()
# ---------- 사용안내 ----------
ws=wb.active; ws.title="사용안내"; ws.sheet_view.showGridLines=False
h(ws,'B2',"27SS 소재일정 역산 트래커 — 출고일 기준, 설·춘절 보정 + 비수기 선정 + 금액 요약 (0808 원본)",16)
sec=[("정렬 순서",BRAND,[
 "① S코드(봄) 전체 — 납기상 최우선 그룹",
 "② 비수기 선정 스타일(봄 제외) — 봄코드 바로 아래 배치 (출고일과 무관하게 우선 조율 대상)",
 "③ 나머지 M(여름) → X(사계절) — 각 그룹 안에서는 출고일 주차 순"]),
("계산 로직",BRAND,[
 "① 출고예정일 = 출고일 주차를 날짜로 환산 (1주=1일, 2주=8일, 3주=15일, 4주=22일 / 9~12월은 전년도)",
 "② 제품입고 시점 = 출고예정일 − 2주(14일). 입고시점이 한국 설 연휴(2/5~2/8)에 걸치면 연휴 전(2/4)으로 앞당김",
 "③ 소재입고 마감일 = 제품입고 시점 − 90일 − 공장휴무 보정일 (중국·베트남 춘절 2/1~2/15, 한국 설 2/5~2/8)",
 "④ 작지 미발행(X/부분) 또는 소재입고일 > 마감일 → S코드(봄) '★ 납기위험', M/X 시즌 '◆ 지연주의'",
 "⑤ 소재입고 책임: CMT = 본사(소재팀) / 완사입·ODM = 생산처 (수/C/완 컬럼 기준 자동)",
 "⑥ 원가합 = Σ(기획수량×입고원가), 소매가합 = Σ(기획수량×예상소매가) — 0808 원본 총원가/총공급가 컬럼 기준"]),
("금액 요약 시트",BRAND,[
 "· 구분: 의류(J/T 등) / 용품(Q로 시작) / WM·피싱(8로 시작) — 품번 첫 글자 기준 자동 분류",
 "· 구분별 스타일 수, 기획수량, 원가합, 소매가합 집계 + 의류 내 비수기 선정 금액 비중 표시",
 "· 트래커에서 수량·금액·비수기(O) 값을 수정하면 요약도 자동 갱신 (SUMPRODUCT 수식)"]),
("비수기 선정",BRAND,[
 "· 'KS_비수기 선정.xlsx' > 27SS비수기 탭 > N열(비수기 1차 확정) 'O' 표시 스타일 = 비수기 선정",
 "· 품번 매칭: B열(품번+컬러) 기준, 스타일번호 변경분은 A열(기존 스타일넘버)로 교차 매칭 (49/50 매칭)",
 "· JWTBM27901→JWTBM27165(변경품번)는 0808 원본에 없어 미반영, 확인 필요"]),
("입력 안내 (노란색 셀 = 입력, 검정 = 자동계산)",BRAND,[
 "· 출고일(주차) 수정 시 입고시점·소재마감·판정 자동 갱신 / 작지발행·사양확정·비수기(O/X) 수정 가능",
 "· 소재입고 예정/실제일, 조율상태, 메모를 입력하세요."]),
("데이터 출처",RED,[
 "· 원본: 0808 첨부 파일 '원본_0808' 시트 (품번+컬러 1,190행 → 424 스타일 집계, 2026-08-08 기준)",
 "· 작지발행: 0808 원본 '작지발행' 건수 / 컬러 구성: 컬러코드 + 브랜드 컬러명 매핑 / 사양확정: 전산 EDW (8/6, 정보성)",
 "· 비수기: KS_비수기 선정.xlsx 27SS비수기 탭 (8/6 수정본) N열 'O' 기준",
 "· 원단처/겉감/안감/생산처(소재DB): '27SS_소재DB' Merged Data 탭과 IMPORTRANGE 실시간 연동 (숨김 탭 '소재DB_RAW')",
 "· 최초 1회: 시트를 열고 '소재DB_RAW' 탭(숨김 해제)에서 A3 셀의 '액세스 허용' 클릭 필요 — 이후 자동 갱신",
 "· 생산처(소재DB) 열은 크로스체크용 — 0808 생산처와 다르면 빨간색 표시 (픽스처/픽스쳐 같은 표기 차이도 포함되니 확인 필요)",
 "· WM/피싱(8로 시작)은 0808 원본에 기획수량 0으로 입력되어 수량 합계 공란 — 금액은 원본 총원가/총공급가 기준",
 "· 8UTCM27101은 컬러별 출고일 상이(1월3주/4월2주) → 다수 값 적용, 확인 필요"])]
r=4
for title,color,items in sec:
    ws[f'B{r}']=title; ws[f'B{r}'].font=Font(name=F,bold=True,size=12,color=color); r+=1
    for t in items:
        ws[f'B{r}']=t; ws[f'B{r}'].font=Font(name=F,size=11 if color==BRAND else 10,color="212121" if color==BRAND else "666666"); r+=1
    r+=1
ws.column_dimensions['B'].width=114

# ---------- 기준정보 ----------
ws=wb.create_sheet("기준정보"); ws.sheet_view.showGridLines=False
h(ws,'B2',"기준정보 (노란색 = 조정 가능한 입력값)",13)
D=datetime.datetime
labels=[("기준일 (오늘)","=TODAY()","판정 기준일. 파일 열 때 자동 갱신"),
        ("입고 리드타임(일)",14,"출고예정일 − 2주 = 제품입고 시점"),
        ("소재 리드타임(일)",90,"제품입고 시점 − 90일 = 소재입고 마감일"),
        ("시즌 연도",2027,"27SS 출고 연도. 9~12월 주차는 전년(2026)으로 환산"),
        ("춘절 공장휴무 시작",D(2027,2,1),"중국·베트남 공장 가동중단 가정 (전후 2주)"),
        ("춘절 공장휴무 종료",D(2027,2,15),"실제 공장별 휴무 확정 시 조정"),
        ("한국 설 연휴 시작",D(2027,2,5),"설날 2/6(토), 연휴 2/5(금)~2/7(일)"),
        ("한국 설 연휴 종료",D(2027,2,8),"대체공휴일 2/8(월) 포함")]
for i,(lab,val,note) in enumerate(labels,start=4):
    ws[f'B{i}']=lab; ws[f'B{i}'].font=Font(name=F,bold=True)
    ws[f'C{i}']=val; ws[f'C{i}'].font=Font(name=F,color="0000FF",bold=True)
    ws[f'C{i}'].fill=PatternFill('solid',fgColor=YELLOW); ws[f'C{i}'].border=border
    if isinstance(val,D) or i==4: ws[f'C{i}'].number_format='yyyy-mm-dd'
    ws[f'D{i}']=note; ws[f'D{i}'].font=Font(name=F,size=9,color="666666")
h(ws,'B13',"소재입고 책임 구분 (수/C/완 기준)",12)
ws['B14']="C / CMT"; ws['C14']="본사(소재팀) 책임 — 소재 발주·입고 본사 컨트롤"
ws['B15']="완(완사입) / ODM"; ws['C15']="생산처 책임 — 생산처가 소재 수배·입고 컨트롤"
for rr in (14,15): ws[f'B{rr}'].font=Font(name=F,bold=True,color=BRAND); ws[f'C{rr}'].font=Font(name=F)
for col,wd in [('B',24),('C',52),('D',50)]: ws.column_dimensions[col].width=wd

# ---------- 소재역산 트래커 ----------
ws=wb.create_sheet("소재역산 트래커")
headers=["NO","MD","DS","Style Code","품명","컬러 구성","시즌","생산형태 (수/C/완)","출고차순","출고일(주차)","출고예정일",
         "제품입고 시점 (출고-2주, 설연휴 보정)","공장휴무 보정 (일)","소재입고 마감일 (입고-90일-보정)","마감까지 D-day",
         "비수기 선정","작지발행","사양확정 (전산 8/6)","소재입고 예정/실제일","소재입고 책임","생산처","생산처 (소재DB)","원산지",
         "원단처 (에이전시)","겉감정보 (원단처+품명)","안감정보 (원단처+품명)",
         "기획수량","원가합 (원)","소매가합 (원)","납기1차 (전산참고)","리스크 사유","판정","조율상태","메모/조치사항"]
INPUT={'J','P','Q','R','S','AG','AH'}
ws.freeze_panes='F5'
h(ws,'A1',"27SS 소재일정 역산 트래커 — 정렬: S코드(봄) → 비수기 선정 → 나머지 (424 스타일)",14)
ws['A2']="기준일:"; ws['A2'].font=Font(name=F,size=10,color="666666")
ws['B2']="=기준정보!$C$4"; ws['B2'].number_format='yyyy-mm-dd'; ws['B2'].font=Font(name=F,size=10,bold=True)
ws['C2']="리드타임:"; ws['C2'].font=Font(name=F,size=10,color="666666")
ws['D2']='="입고 "&기준정보!$C$5&"일 + 소재 "&기준정보!$C$6&"일 + 휴무보정"'; ws['D2'].font=Font(name=F,size=10,bold=True)
ws['G2']="노란색 = 입력 · ★ = S코드 작지 미발행/소재마감 초과 · 정렬: ①봄(S) ②비수기 선정 ③나머지 · 금액은 [금액 요약] 참조"
ws['G2'].font=Font(name=F,size=10,color=RED)
HR=4
for ci,hd in enumerate(headers,start=1):
    c=ws.cell(row=HR,column=ci,value=hd)
    c.font=Font(name=F,size=10,bold=True,color="FFFFFF"); c.fill=PatternFill('solid',fgColor=BRAND)
    c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=border

def week_expr(x):
    return (f'IFERROR(DATE(IF(VALUE(LEFT({x},FIND("월",{x})-1))>=9,기준정보!$C$7-1,기준정보!$C$7),'
            f'VALUE(LEFT({x},FIND("월",{x})-1)),1+7*(VALUE(SUBSTITUTE(MID({x},FIND("월",{x})+1,5),"주",""))-1)),"")')

N=len(rows)
for i,r in enumerate(rows):
    rw=HR+1+i
    nap = r['납기1차(참고)']
    nap_v = datetime.datetime.strptime(nap,'%Y-%m-%d') if re.match(r'^\d{4}-\d{2}-\d{2}$',nap or '') else (nap or None)
    L_f=(f'=IF($K{rw}="","",IF(AND($K{rw}-기준정보!$C$5>=기준정보!$C$10,$K{rw}-기준정보!$C$5<=기준정보!$C$11),'
         f'기준정보!$C$10-1,$K{rw}-기준정보!$C$5))')
    M_f=(f'=IF($L{rw}="","",IF(OR(ISNUMBER(SEARCH("중국",$W{rw})),ISNUMBER(SEARCH("베트남",$W{rw}))),'
         f'MAX(0,MIN($L{rw},기준정보!$C$9)-MAX($L{rw}-기준정보!$C$6,기준정보!$C$8)+1),'
         f'IF(ISNUMBER(SEARCH("한국",$W{rw})),MAX(0,MIN($L{rw},기준정보!$C$11)-MAX($L{rw}-기준정보!$C$6,기준정보!$C$10)+1),0)))')
    vals={
     'A':i+1,'B':r['MD'],'C':r['DS'],'D':r['StyleCode'],'E':r['품명'],'F':r['컬러구성'],
     'G':f'=IF($D{rw}="","",IF(MID($D{rw},5,1)="S","S(봄)",IF(MID($D{rw},5,1)="M","M(여름)",IF(MID($D{rw},5,1)="X","X(사계절)",MID($D{rw},5,1)))))',
     'H':r['생산형태'],'I':r['출고차순'],'J':r['출고일(주차)'] or None,
     'K':f'=IF($J{rw}="","",{week_expr(f"$J{rw}")})',
     'L':L_f,'M':M_f,
     'N':f'=IF($L{rw}="","",$L{rw}-기준정보!$C$6-$M{rw})',
     'O':f'=IF($N{rw}="","",$N{rw}-기준정보!$C$4)',
     'P':r['비수기'] or None,
     'Q':r['작지발행'],'R':r['사양확정'] or None,'S':None,
     'T':(f'=IF($H{rw}="","",IF(AND(ISNUMBER(SEARCH("CMT",$H{rw})),OR(ISNUMBER(SEARCH("완사입",$H{rw})),ISNUMBER(SEARCH("ODM",$H{rw})))),'
          f'"본사/생산처 혼재",IF(ISNUMBER(SEARCH("CMT",$H{rw})),"본사(소재팀)","생산처")))'),
     'U':r['생산처'],
     'V':f'=IFERROR(TEXTJOIN(" / ",TRUE,UNIQUE(FILTER(소재DB_RAW!$K$3:$K,소재DB_RAW!$G$3:$G=$D{rw}))),"")',
     'W':r['원산지'],
     'X':f'=IFERROR(TEXTJOIN(" / ",TRUE,UNIQUE(FILTER(소재DB_RAW!$N$3:$N,소재DB_RAW!$G$3:$G=$D{rw}))),"")',
     'Y':(f'=IFERROR(TEXTJOIN(CHAR(10),TRUE,UNIQUE(FILTER(소재DB_RAW!$N$3:$N&") "&소재DB_RAW!$O$3:$O,'
          f'(소재DB_RAW!$G$3:$G=$D{rw})*ISNUMBER(SEARCH("겉감",소재DB_RAW!$M$3:$M))))),"")'),
     'Z':(f'=IFERROR(TEXTJOIN(CHAR(10),TRUE,UNIQUE(FILTER(소재DB_RAW!$N$3:$N&") "&소재DB_RAW!$O$3:$O,'
          f'(소재DB_RAW!$G$3:$G=$D{rw})*ISNUMBER(SEARCH("안감",소재DB_RAW!$M$3:$M))))),"")'),
     'AA':int(r['기획수량'] or 0),
     'AB':int(r['원가합'] or 0),'AC':int(r['소매가합'] or 0),'AD':nap_v,
     'AE':(f'=IF($K{rw}="","출고일 미정",IF($Q{rw}="X","작지 미발행",IF($Q{rw}="부분","작지 일부 미발행",'
          f'IF(AND($S{rw}<>"",$S{rw}>$N{rw}),"소재입고 마감("&TEXT($N{rw},"MM/DD")&") 초과",'
          f'IF($Q{rw}="","발행여부 확인 필요",IF($S{rw}="","소재 입고일정 확인 필요","정상"))))))'),
     'AF':(f'=IF($K{rw}="","",IF(OR(LEFT($AE{rw},2)="작지",LEFT($AE{rw},4)="소재입고"),'
          f'IF(LEFT($G{rw},1)="S","★ 납기위험(S코드)","◆ 지연주의"),IF($AE{rw}="정상","✓ 정상","? 확인필요")))'),
     'AG':None,'AH':None,
    }
    for col,v in vals.items():
        c=ws[f'{col}{rw}']
        if v is not None: c.value=v
        c.font=Font(name=F,size=10); c.border=border
        if col in INPUT: c.fill=PatternFill('solid',fgColor="FFF9C4")
        if col in ('K','L','N','S') or (col=='AD' and isinstance(v,datetime.datetime)): c.number_format='yyyy-mm-dd'
        if col=='O': c.number_format='0;[RED]-0'
        if col=='M': c.number_format='0'
        if col in ('AA','AB','AC'): c.number_format='#,##0'
        if col in ('Y','Z'): c.alignment=Alignment(wrap_text=True, vertical='top')
        if col in ('A','G','I','J','M','O','P','Q','R','AF','AG'): c.alignment=Alignment(horizontal='center')

last=HR+N
for f1,rng in [('"O,X"',f'P{HR+1}:P{last}'),('"O,부분,X"',f'Q{HR+1}:Q{last}'),('"O,부분,X"',f'R{HR+1}:R{last}'),('"미조율,조율중,조율완료"',f'AG{HR+1}:AG{last}')]:
    dv=DataValidation(type="list",formula1=f1,allow_blank=True); ws.add_data_validation(dv); dv.add(rng)
def cf(col,txt,color,fc):
    dxf=DifferentialStyle(fill=PatternFill(start_color=color,end_color=color,fill_type='solid'),font=Font(name=F,color=fc,bold=True))
    rule=Rule(type="containsText",operator="containsText",text=txt,dxf=dxf)
    rule.formula=[f'NOT(ISERROR(SEARCH("{txt}",{col}{HR+1})))']
    ws.conditional_formatting.add(f'{col}{HR+1}:{col}{last}',rule)
cf('AF',"★","F8CBCC",RED); cf('AF',"◆","FDE9CC","9C5700"); cf('AF',"✓",TOTALG,BRAND); cf('AF',"?","EEEEEE","666666")
# 생산처 크로스체크: 소재DB 생산처에 0808 생산처가 포함되지 않으면 강조
dxf_mm=DifferentialStyle(fill=PatternFill(start_color="FDE3E4",end_color="FDE3E4",fill_type='solid'),font=Font(name=F,color=RED,bold=True))
rule_mm=Rule(type="expression",dxf=dxf_mm)
rule_mm.formula=[f'AND($V{HR+1}<>"",$U{HR+1}<>"",ISERROR(SEARCH($U{HR+1},$V{HR+1})))']
ws.conditional_formatting.add(f'V{HR+1}:V{last}',rule_mm)
cf('P',"O","DCEDF7",NAVY)
widths={'A':5,'B':9,'C':9,'D':13,'E':30,'F':24,'G':9,'H':13,'I':9,'J':12,'K':12,'L':14,'M':9,'N':14,'O':9,'P':9,'Q':9,'R':11,'S':13,'T':13,'U':12,'V':13,'W':10,'X':16,'Y':34,'Z':30,'AA':9,'AB':14,'AC':15,'AD':13,'AE':26,'AF':15,'AG':9,'AH':24}
for col,wd in widths.items(): ws.column_dimensions[col].width=wd
ws.auto_filter.ref=f'A{HR}:AH{last}'

# ---------- 소재DB_RAW (실시간 연동 헬퍼) ----------
ws=wb.create_sheet("소재DB_RAW")
ws['A1']="27SS_소재DB > Merged Data 실시간 연동 (IMPORTRANGE). 최초 1회 A3 셀에서 '액세스 허용'을 클릭해야 합니다. 이 탭은 수정하지 마세요."
ws['A1'].font=Font(name=F,size=10,bold=True,color=RED)
ws['A3']='=IMPORTRANGE("https://docs.google.com/spreadsheets/d/1_LNPz9XGRmRAJ9Y4cqtFF8NtZTLSIld4vtzLJlOcka0","Merged Data!A3:P2000")'
ws.sheet_state='hidden'

TRK="'소재역산 트래커'"
RNG_D=f'{TRK}!$D$5:$D${last}'; RNG_W=f'{TRK}!$AA$5:$AA${last}'; RNG_X=f'{TRK}!$AB$5:$AB${last}'; RNG_Y=f'{TRK}!$AC$5:$AC${last}'; RNG_P=f'{TRK}!$P$5:$P${last}'

# ---------- 금액 요약 ----------
ws=wb.create_sheet("금액 요약"); ws.sheet_view.showGridLines=False
h(ws,'B2',"27SS 구분별 금액 요약 — 원가합 · 소매가합",14)
ws['B3']="구분: 의류(J/T 등) · 용품(품번 Q로 시작) · WM/피싱(품번 8로 시작). 트래커 값 수정 시 자동 갱신됩니다."
ws['B3'].font=Font(name=F,size=10,color="666666")
hdrs=["구분","스타일 수","기획수량 (장)","원가합 (원)","소매가합 (원)","소매가합 비중"]
for ci,t in enumerate(hdrs,start=2):
    c=ws.cell(row=5,column=ci,value=t)
    c.font=Font(name=F,size=10,bold=True,color="FFFFFF"); c.fill=PatternFill('solid',fgColor=BRAND)
    c.alignment=Alignment(horizontal='center'); c.border=border
Q='LEFT('+RNG_D+',1)="Q"'; E8='LEFT('+RNG_D+',1)="8"'
NOTQ8=f'(LEFT({RNG_D},1)<>"Q")*(LEFT({RNG_D},1)<>"8")*({RNG_D}<>"")'
cats=[("의류",NOTQ8),("용품 (Q)",f'({Q})*1'),("WM/피싱 (8)",f'({E8})*1')]
for ri,(name,cond) in enumerate(cats,start=6):
    ws[f'B{ri}']=name
    ws[f'C{ri}']=f'=SUMPRODUCT({cond})'
    ws[f'D{ri}']=f'=SUMPRODUCT({cond}*{RNG_W})'
    ws[f'E{ri}']=f'=SUMPRODUCT({cond}*{RNG_X})'
    ws[f'F{ri}']=f'=SUMPRODUCT({cond}*{RNG_Y})'
    ws[f'G{ri}']=f'=IF($F$9=0,"",F{ri}/$F$9)'
tot=9
ws[f'B{tot}']="합계"
ws[f'C{tot}']='=SUM(C6:C8)'; ws[f'D{tot}']='=SUM(D6:D8)'; ws[f'E{tot}']='=SUM(E6:E8)'; ws[f'F{tot}']='=SUM(F6:F8)'; ws[f'G{tot}']='=IF(F9=0,"",1)'
for ri in range(6,10):
    for col in 'BCDEFG':
        c=ws[f'{col}{ri}']; c.font=Font(name=F,size=11,bold=(ri==tot)); c.border=border
        if col in 'CDEF': c.number_format='#,##0'
        if col=='G': c.number_format='0.0%'
        if col!='B': c.alignment=Alignment(horizontal='right')
    if ri==tot:
        for col in 'BCDEFG':
            ws[f'{col}{ri}'].fill=PatternFill('solid',fgColor=TOTALG)
            ws[f'{col}{ri}'].font=Font(name=F,size=11,bold=True,color=BRAND)

h(ws,'B12',"의류 내 비수기 선정 비중",13)
ws['B13']="비수기 = KS_비수기 선정 27SS비수기 탭 N열 'O' 확정 스타일 (트래커 '비수기 선정' 열 기준)"
ws['B13'].font=Font(name=F,size=10,color="666666")
hdrs2=["구분","스타일 수","기획수량 (장)","원가합 (원)","소매가합 (원)"]
for ci,t in enumerate(hdrs2,start=2):
    c=ws.cell(row=15,column=ci,value=t)
    c.font=Font(name=F,size=10,bold=True,color="FFFFFF"); c.fill=PatternFill('solid',fgColor=BRAND)
    c.alignment=Alignment(horizontal='center'); c.border=border
BIS=f'({RNG_P}="O")'
ws['B16']="의류 전체"
ws['C16']='=C6'; ws['D16']='=D6'; ws['E16']='=E6'; ws['F16']='=F6'
ws['B17']="비수기 선정 (의류)"
ws['C17']=f'=SUMPRODUCT({NOTQ8}*{BIS})'
ws['D17']=f'=SUMPRODUCT({NOTQ8}*{BIS}*{RNG_W})'
ws['E17']=f'=SUMPRODUCT({NOTQ8}*{BIS}*{RNG_X})'
ws['F17']=f'=SUMPRODUCT({NOTQ8}*{BIS}*{RNG_Y})'
ws['B18']="비수기 비중"
for col in 'CDEF':
    ws[f'{col}18']=f'=IF({col}16=0,"",{col}17/{col}16)'
for ri in range(16,19):
    for col in 'BCDEF':
        c=ws[f'{col}{ri}']; c.font=Font(name=F,size=11); c.border=border
        if ri<18 and col in 'CDEF': c.number_format='#,##0'
        if ri==18 and col in 'CDEF': c.number_format='0.0%'
        if col!='B': c.alignment=Alignment(horizontal='right')
for col in 'BCDEF':
    ws[f'{col}18'].fill=PatternFill('solid',fgColor="DCEDF7"); ws[f'{col}18'].font=Font(name=F,size=11,bold=True,color=NAVY)
ws['B20']="* 원가합 = Σ(기획수량×입고원가), 소매가합 = Σ(기획수량×예상소매가) — 0808 원본 총원가·총공급가 기준"
ws['B20'].font=Font(name=F,size=9,color="666666")
for col,wd in [('B',20),('C',11),('D',14),('E',18),('F',19),('G',13)]: ws.column_dimensions[col].width=wd

# ---------- 리스크 요약 ----------
ws=wb.create_sheet("리스크 요약"); ws.sheet_view.showGridLines=False
h(ws,'B2',"27SS 소재일정 리스크 요약 (0808 원본 · 설/춘절 보정 · 비수기 선정)",14)
ws['B3']=(f'="기준일: "&TEXT(기준정보!$C$4,"YYYY-MM-DD")&"  ·  소재마감 = 출고일 − "&(기준정보!$C$5+기준정보!$C$6)&"일 − 휴무보정"'
          f'&"  ·  S "&COUNTIF({TRK}!$G$5:$G${last},"S(봄)")&" / M "&COUNTIF({TRK}!$G$5:$G${last},"M(여름)")&" / X "&COUNTIF({TRK}!$G$5:$G${last},"X(사계절)")')
ws['B3'].font=Font(name=F,size=10,color="666666")
kpis=[("전체 스타일",f'=COUNTA({TRK}!$D$5:$D${last})',BRAND),
      ("S코드(봄) 스타일",f'=COUNTIF({TRK}!$G$5:$G${last},"S(봄)")',BRAND),
      ("★ 납기위험(S코드)",f'=COUNTIF({TRK}!$AF$5:$AF${last},"★*")',RED),
      ("◆ 지연주의(M/X)",f'=COUNTIF({TRK}!$AF$5:$AF${last},"◆*")',ORANGE),
      ("비수기 선정 스타일",f'=COUNTIF({RNG_P},"O")',NAVY)]
for i,(lab,f_,col_) in enumerate(kpis):
    col=get_column_letter(2+i*2)
    ws[f'{col}5']=lab; ws[f'{col}5'].font=Font(name=F,size=10,color="666666")
    ws[f'{col}6']=f_; ws[f'{col}6'].font=Font(name=F,size=22,bold=True,color=col_)
h(ws,'B9',"MD별 현황",12)
hdrs=["MD","스타일 수","S코드","작지 미발행(X/부분)","사양확정 미완(X/부분)","비수기 선정","★ 납기위험","조율완료"]
for ci,t in enumerate(hdrs,start=2):
    c=ws.cell(row=10,column=ci,value=t)
    c.font=Font(name=F,size=10,bold=True,color="FFFFFF"); c.fill=PatternFill('solid',fgColor=BRAND)
    c.alignment=Alignment(horizontal='center'); c.border=border
mds=sorted({r['MD'] for r in rows if r['MD']})
for ri,md in enumerate(mds,start=11):
    ws[f'B{ri}']=md
    ws[f'C{ri}']=f'=COUNTIF({TRK}!$B$5:$B${last},$B{ri})'
    ws[f'D{ri}']=f'=COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$G$5:$G${last},"S(봄)")'
    ws[f'E{ri}']=f'=COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$Q$5:$Q${last},"X")+COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$Q$5:$Q${last},"부분")'
    ws[f'F{ri}']=f'=COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$R$5:$R${last},"X")+COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$R$5:$R${last},"부분")'
    ws[f'G{ri}']=f'=COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{RNG_P},"O")'
    ws[f'H{ri}']=f'=COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$AF$5:$AF${last},"★*")'
    ws[f'I{ri}']=f'=COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$AG$5:$AG${last},"조율완료")'
    for col in 'BCDEFGHI':
        ws[f'{col}{ri}'].font=Font(name=F,size=10); ws[f'{col}{ri}'].border=border
        if col!='B': ws[f'{col}{ri}'].alignment=Alignment(horizontal='center')
tot2=11+len(mds)
ws[f'B{tot2}']="합계"; ws[f'B{tot2}'].font=Font(name=F,size=10,bold=True,color=BRAND)
for ci in range(3,10):
    cl=get_column_letter(ci)
    ws[f'{cl}{tot2}']=f'=SUM({cl}11:{cl}{tot2-1})'
    ws[f'{cl}{tot2}'].font=Font(name=F,size=10,bold=True,color=BRAND); ws[f'{cl}{tot2}'].alignment=Alignment(horizontal='center')
for col in 'BCDEFGHI':
    ws[f'{col}{tot2}'].border=border; ws[f'{col}{tot2}'].fill=PatternFill('solid',fgColor=TOTALG)
msg_r=tot2+3
h(ws,f'B{msg_r}',"공유 메시지 (관계부서 발송용)",12)
msg=("27SS 영업MD 출고계획은 기획MD 출고계획과 거의 일치하게 운영되며, 생산도 동 일정에 맞춰야 합니다. "
     "관건은 소재일정의 정확성입니다. 출고일 2주 전에는 제품이 입고되어야 하고(27년 설 연휴 2/5~2/8 입고 불가 반영), "
     "제품입고 기준 90일 전에는 소재가 입고되어야 합니다. 중국·베트남 공장은 춘절 휴무(2/1~2/15 가정)로 생산기간이 "
     "겹치는 만큼 소재 마감이 추가로 앞당겨집니다. 이 마감까지 소재 입고가 확인되지 않았거나, 현시점 작업지시서가 "
     "발행되지 않은 S코드(봄상품)는 납기 준수가 어렵습니다. 비수기 1차 확정 스타일(의류 소매가 기준 약 28%)은 "
     "BT 컨펌·8월말 원단 도착 일정 준수가 전제이므로 소재 일정 회신 1순위로 요청드립니다. "
     "CMT는 본사(소재팀), 완사입·ODM은 생산처에서 소재 입고 가능일을 회신해 주시기 바랍니다.")
ws.merge_cells(f'B{msg_r+1}:I{msg_r+6}')
c=ws[f'B{msg_r+1}']; c.value=msg
c.font=Font(name=F,size=11); c.alignment=Alignment(wrap_text=True,vertical='top'); c.fill=PatternFill('solid',fgColor=PAPER)
for col,wd in [('B',16),('C',12),('D',12),('E',18),('F',19),('G',12),('H',13),('I',12)]: ws.column_dimensions[col].width=wd

wb.save('output/27SS_소재일정_역산_트래커.xlsx')
print("saved. styles:", N, "last:", last)
