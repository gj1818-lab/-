# 27FW 소재일정 역산 트래커 빌드 스크립트 (27SS build_tracker.py 골격 재사용)
# 원본: data/style_master_27fw.csv (생진테 Total(GR,YL통합) 추출 집계본)
# 27FW 확정 기준: 연도 매핑 전부 2027 / 출고-14일=제품입고, 입고-100일=소재마감 /
#   휴무 보정 = 미얀마 띤쨘(4/9~4/18)·인도네시아 르바란(3/4~3/17), 원산지 국가별 적용
# 정렬: ① A코드(가을) → ② 비수기 선정(가을 제외, 데이터 대기) → ③ W(겨울)→X(사계절), 그룹 내 출고주차 순
import csv, re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

BRAND="2C5F2D"; PAPER="F5F5F5"; RED="E8474C"; ORANGE="F18F01"; TOTALG="E8F5E9"; YELLOW="FFFF00"; NAVY="1B3A4B"
F="맑은 고딕"
thin=Side(style='thin', color='D9D9D9'); border=Border(left=thin,right=thin,top=thin,bottom=thin)

def h(ws,cell,text,size=11,bold=True,color="FFFFFF",fill=BRAND):
    ws[cell]=text; ws[cell].font=Font(name=F,size=size,bold=bold,color=color)
    if fill: ws[cell].fill=PatternFill('solid',fgColor=fill)

rows=list(csv.DictReader(open('data/style_master_27fw.csv')))

def week_key(w):
    m=re.match(r'(\d+)월(\d+)주',w or '')
    return (int(m.group(1)),int(m.group(2))) if m else (99,99)  # 전부 2027년 — 월 순서 그대로

def sort_key(r):
    if r['시즌코드']=='A': grp=0
    elif r.get('비수기','')=='O': grp=1
    else: grp=2
    season_order={'W':0,'X':1}.get(r['시즌코드'],2) if grp==2 else 0
    return (grp, season_order, week_key(r['출고일(주차)']), r['StyleCode'])
rows.sort(key=sort_key)

def money(v):
    v=(v or '').replace(',','').strip()
    return int(v) if v.isdigit() else 0

wb=Workbook()
# ---------- 사용안내 ----------
ws=wb.active; ws.title="사용안내"; ws.sheet_view.showGridLines=False
h(ws,'B2',"27FW 소재일정 역산 트래커 — 생진테 출고일 기준, 띤쨘·르바란 보정 (v1)",16)
sec=[("정렬 순서",BRAND,[
 "① A코드(가을) 전체 — 납기상 최우선 그룹",
 "② 비수기 선정 스타일(가을 제외) — 27FW 비수기 선정 수령 시 반영 예정 (현재 공란)",
 "③ 나머지 W(겨울) → X(사계절) — 각 그룹 안에서는 출고일 주차 순"]),
("계산 로직",BRAND,[
 "① 출고예정일 = 생진테 출고일 주차를 날짜로 환산 (1주=1일, 2주=8일, 3주=15일, 4주=22일 / 전 주차 2027년)",
 "② 제품입고 시점 = 출고예정일 − 2주(14일). 입고시점이 생산국 휴무에 걸치면 휴무 전으로 앞당김",
 "③ 소재입고 마감일 = 제품입고 시점 − 100일 − 공장휴무 보정일 (FW 난이도 반영, 27SS는 90일)",
 "④ 공장휴무 보정: 생산기간(마감~입고)이 휴무와 겹친 일수만큼 마감 추가 앞당김 —",
 "   미얀마 생산분 = 띤쨘 4/9~4/18(10일), 인도네시아 생산분 = 르바란 3/4~3/17(14일). 기준정보에서 조정 가능",
 "⑤ 작지 미발행(X/부분) 또는 소재입고일 > 마감일 → A코드(가을) '★ 납기위험', W/X '◆ 지연주의'",
 "⑥ 소재입고 책임: C(CMT) = 본사(소재팀) / 완(완사입) = 생산처 (수/C/완 기준 자동)"]),
("입력 안내 (노란색 셀 = 입력, 검정 = 자동계산)",BRAND,[
 "· 출고일(주차) 수정 시 입고시점·소재마감·판정 자동 갱신 (형식: 8월1주)",
 "· 비수기(O)·작지발행·사양확정(O/부분/X)은 데이터 수령 시 채움 — 수동 입력도 가능",
 "· 소재입고 예정/실제일, 조율상태, 메모를 입력하세요."]),
("데이터 출처",RED,[
 "· 원본: (자동화)_KS_2027FW_GREEN Product Line Plan (생진테/리퀘스트 통합) Total(GR,YL통합) 탭",
 "  → 27FW_생진테_추출 시트(IMPORTRANGE) 경유, 2026-09-16 기준 853 컬러행 → 425 스타일 집계",
 "· 출고일 미입력 117 스타일은 '출고일 미정' 표시 — 생진테에 주차 입력되면 추출 시트 자동 갱신 (트래커는 재빌드 또는 수동 입력)",
 "· 비수기 선정(27FW)·작지발행/사양확정(EDW)·소재DB(원단처/겉감/안감/LAB DIP)는 원본 수령 대기 — 해당 열 공란",
 "· 이월 품번 1건 포함: QYHFW25421 (25년 기획, 생진테 등재분)"])]
r=4
for title,color,items in sec:
    ws[f'B{r}']=title; ws[f'B{r}'].font=Font(name=F,bold=True,size=12,color=color); r+=1
    for t in items:
        ws[f'B{r}']=t; ws[f'B{r}'].font=Font(name=F,size=11 if color==BRAND else 10,color="212121" if color==BRAND else "666666"); r+=1
    r+=1
ws.column_dimensions['B'].width=114

# ---------- 기준정보 ----------
ws=wb.create_sheet("기준정보"); ws.sheet_view.showGridLines=False
h(ws,'B2',"기준정보 (노란색 = 조정 가능한 입력값)",13)
D=datetime.datetime
labels=[("기준일 (오늘)","=TODAY()","판정 기준일. 파일 열 때 자동 갱신"),          # C4
        ("입고 리드타임(일)",14,"출고예정일 − 2주 = 제품입고 시점"),               # C5
        ("소재 리드타임(일)",100,"제품입고 시점 − 100일 = 소재입고 마감일 (FW 난이도 반영)"),  # C6
        ("시즌 연도",2027,"27FW 출고 주차는 전부 2027년으로 환산 (담당자 확정)"),   # C7
        ("띤쨘 휴무 시작",D(2027,4,9),"미얀마 생산분. 공휴일 4/13~17 + 앞뒤 생산처 추가 휴무"),  # C8
        ("띤쨘 휴무 종료",D(2027,4,18),"실제 공장별 휴무 확정 시 조정"),            # C9
        ("르바란 휴무 시작",D(2027,3,4),"인도네시아 생산분. Eid 3/10 전후 귀향(mudik) 포함 2주"),  # C10
        ("르바란 휴무 종료",D(2027,3,17),"실제 공장별 휴무 확정 시 조정")]          # C11
for i,(lab,val,note) in enumerate(labels,start=4):
    ws[f'B{i}']=lab; ws[f'B{i}'].font=Font(name=F,bold=True)
    ws[f'C{i}']=val; ws[f'C{i}'].font=Font(name=F,color="0000FF",bold=True)
    ws[f'C{i}'].fill=PatternFill('solid',fgColor=YELLOW); ws[f'C{i}'].border=border
    if isinstance(val,D) or i==4: ws[f'C{i}'].number_format='yyyy-mm-dd'
    ws[f'D{i}']=note; ws[f'D{i}'].font=Font(name=F,size=9,color="666666")
h(ws,'B13',"소재입고 책임 구분 (수/C/완 기준)",12)
ws['B14']="C / CMT"; ws['C14']="본사(소재팀) 책임 — 소재 발주·입고 본사 컨트롤"
ws['B15']="완(완사입) / ODM"; ws['C15']="생산처 책임 — 생산처가 소재 수배·입고 컨트롤"
for rr in (14,15): ws[f'B{rr}'].font=Font(name=F,bold=True,color=BRAND); ws[f'C{rr}'].font=Font(name=F)
for col,wd in [('B',24),('C',52),('D',56)]: ws.column_dimensions[col].width=wd

# ---------- 소재역산 트래커 ----------
ws=wb.create_sheet("소재역산 트래커")
headers=["NO","MD","DS","Style Code","품명","컬러 구성","시즌","생산형태 (수/C/완)","출고일(주차)","출고예정일",
         "제품입고 시점 (출고-2주, 휴무 보정)","공장휴무 보정 (일)","소재입고 마감일 (입고-100일-보정)","마감까지 D-day",
         "비수기 선정 (대기)","작지발행 (일매·EDW 대기)","사양확정 (EDW 대기)","소재입고 예정/실제일","소재입고 책임","생산처","원산지",
         "원단처 (소재DB 대기)","겉감정보 (대기)","안감정보 (대기)","LAB DIP CFM (대기)",
         "기획수량","원가합 (원)","소매가합 (원)","입고예정일 (생진테 참고)","리스크 사유","판정","조율상태","메모/조치사항"]
INPUT={'I','O','P','Q','R','AF','AG'}
ws.freeze_panes='F5'
h(ws,'A1',f"27FW 소재일정 역산 트래커 — 정렬: A코드(가을) → 비수기 선정 → W→X ({len(rows)} 스타일)",14)
ws['A2']="기준일:"; ws['A2'].font=Font(name=F,size=10,color="666666")
ws['B2']="=기준정보!$C$4"; ws['B2'].number_format='yyyy-mm-dd'; ws['B2'].font=Font(name=F,size=10,bold=True)
ws['C2']="리드타임:"; ws['C2'].font=Font(name=F,size=10,color="666666")
ws['D2']='="입고 "&기준정보!$C$5&"일 + 소재 "&기준정보!$C$6&"일 + 휴무보정"'; ws['D2'].font=Font(name=F,size=10,bold=True)
ws['G2']="노란색 = 입력 · ★ = A코드 작지 미발행/소재마감 초과 · 비수기/작지/사양확정/소재DB 열은 원본 수령 대기"
ws['G2'].font=Font(name=F,size=10,color=RED)
HR=4
for ci,hd in enumerate(headers,start=1):
    c=ws.cell(row=HR,column=ci,value=hd)
    c.font=Font(name=F,size=10,bold=True,color="FFFFFF"); c.fill=PatternFill('solid',fgColor=BRAND)
    c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=border

def week_expr(x):
    # 전부 2027년(기준정보 C7) — 27SS의 9~12월 전년 규칙 제거
    return (f'IFERROR(DATE(기준정보!$C$7,VALUE(LEFT({x},FIND("월",{x})-1)),'
            f'1+7*(VALUE(SUBSTITUTE(MID({x},FIND("월",{x})+1,5),"주",""))-1)),"")')

N=len(rows)
for i,r in enumerate(rows):
    rw=HR+1+i
    qty=money(r['기획수량']); cost=money(r['입고원가']); retail=money(r['예상소매가'])
    ipd=r['입고예정일']
    ipd_v=datetime.datetime.strptime(ipd,'%Y-%m-%d') if re.match(r'^\d{4}-\d{2}-\d{2}$',ipd or '') else (ipd or None)
    # K: 제품입고 = 출고 − 입고리드, 생산국 휴무에 걸치면 휴무 전일로 앞당김 (미얀마=띤쨘, 인니=르바란)
    K_f=(f'=IF($J{rw}="","",IF(AND(ISNUMBER(SEARCH("미얀마",$U{rw})),$J{rw}-기준정보!$C$5>=기준정보!$C$8,$J{rw}-기준정보!$C$5<=기준정보!$C$9),기준정보!$C$8-1,'
         f'IF(AND(ISNUMBER(SEARCH("인도네시아",$U{rw})),$J{rw}-기준정보!$C$5>=기준정보!$C$10,$J{rw}-기준정보!$C$5<=기준정보!$C$11),기준정보!$C$10-1,$J{rw}-기준정보!$C$5)))')
    # L: 생산기간(입고−100일 ~ 입고)이 생산국 휴무와 겹친 일수
    L_f=(f'=IF($K{rw}="","",IF(ISNUMBER(SEARCH("미얀마",$U{rw})),'
         f'MAX(0,MIN($K{rw},기준정보!$C$9)-MAX($K{rw}-기준정보!$C$6,기준정보!$C$8)+1),'
         f'IF(ISNUMBER(SEARCH("인도네시아",$U{rw})),MAX(0,MIN($K{rw},기준정보!$C$11)-MAX($K{rw}-기준정보!$C$6,기준정보!$C$10)+1),0)))')
    vals={
     'A':i+1,'B':r['MD'],'C':r['DS'],'D':r['StyleCode'],'E':r['품명'],'F':r['컬러구성'],
     'G':f'=IF($D{rw}="","",IF(MID($D{rw},5,1)="A","A(가을)",IF(MID($D{rw},5,1)="W","W(겨울)",IF(MID($D{rw},5,1)="X","X(사계절)",MID($D{rw},5,1)))))',
     'H':r['수/C/완'] or None,'I':r['출고일(주차)'] or None,
     'J':f'=IF($I{rw}="","",{week_expr(f"$I{rw}")})',
     'K':K_f,'L':L_f,
     'M':f'=IF($K{rw}="","",$K{rw}-기준정보!$C$6-$L{rw})',
     'N':f'=IF($M{rw}="","",$M{rw}-기준정보!$C$4)',
     'O':None,'P':None,'Q':None,'R':None,
     'S':f'=IF($H{rw}="","",IF($H{rw}="C","본사(소재팀)",IF($H{rw}="완","생산처","확인 필요")))',
     'T':r['생산처'] or None,'U':r['원산지'] or None,
     'V':None,'W':None,'X':None,'Y':None,
     'Z':qty or None,'AA':qty*cost or None,'AB':qty*retail or None,'AC':ipd_v,
     'AD':(f'=IF($J{rw}="","출고일 미정",IF($P{rw}="X","작지 미발행",IF($P{rw}="부분","작지 일부 미발행",'
          f'IF(AND($R{rw}<>"",$R{rw}>$M{rw}),"소재입고 마감("&TEXT($M{rw},"MM/DD")&") 초과",'
          f'IF($P{rw}="","발행여부 확인 필요",IF($R{rw}="","소재 입고일정 확인 필요","정상"))))))'),
     'AE':(f'=IF($J{rw}="","",IF(OR(LEFT($AD{rw},2)="작지",LEFT($AD{rw},4)="소재입고"),'
          f'IF(LEFT($G{rw},1)="A","★ 납기위험(A코드)","◆ 지연주의"),IF($AD{rw}="정상","✓ 정상","? 확인필요")))'),
     'AF':None,'AG':None,
    }
    for col,v in vals.items():
        c=ws[f'{col}{rw}']
        if v is not None: c.value=v
        c.font=Font(name=F,size=10); c.border=border
        if col in INPUT: c.fill=PatternFill('solid',fgColor="FFF9C4")
        if col in ('J','K','M','R','Y') or (col=='AC' and isinstance(v,datetime.datetime)): c.number_format='yyyy-mm-dd'
        if col=='N': c.number_format='0;[RED]-0'
        if col=='L': c.number_format='0'
        if col in ('Z','AA','AB'): c.number_format='#,##0'
        if col in ('W','X'): c.alignment=Alignment(wrap_text=True, vertical='top')
        if col in ('A','G','H','I','L','N','O','P','Q','AE','AF'): c.alignment=Alignment(horizontal='center')

last=HR+N
for f1,rng in [('"O,X"',f'O{HR+1}:O{last}'),('"O,부분,X"',f'P{HR+1}:P{last}'),('"O,부분,X"',f'Q{HR+1}:Q{last}'),('"미조율,조율중,조율완료"',f'AF{HR+1}:AF{last}')]:
    dv=DataValidation(type="list",formula1=f1,allow_blank=True); ws.add_data_validation(dv); dv.add(rng)
def cf(col,txt,color,fc):
    dxf=DifferentialStyle(fill=PatternFill(start_color=color,end_color=color,fill_type='solid'),font=Font(name=F,color=fc,bold=True))
    rule=Rule(type="containsText",operator="containsText",text=txt,dxf=dxf)
    rule.formula=[f'NOT(ISERROR(SEARCH("{txt}",{col}{HR+1})))']
    ws.conditional_formatting.add(f'{col}{HR+1}:{col}{last}',rule)
cf('AE',"★","F8CBCC",RED); cf('AE',"◆","FDE9CC","9C5700"); cf('AE',"✓",TOTALG,BRAND); cf('AE',"?","EEEEEE","666666")
cf('O',"O","DCEDF7",NAVY)
widths={'A':5,'B':9,'C':9,'D':13,'E':30,'F':22,'G':10,'H':13,'I':12,'J':12,'K':14,'L':9,'M':14,'N':9,'O':9,'P':11,'Q':11,'R':13,'S':13,'T':12,'U':10,'V':14,'W':22,'X':22,'Y':13,'Z':9,'AA':14,'AB':15,'AC':13,'AD':26,'AE':16,'AF':9,'AG':24}
for col,wd in widths.items(): ws.column_dimensions[col].width=wd
ws.auto_filter.ref=f'A{HR}:AG{last}'

TRK="'소재역산 트래커'"
RNG_D=f'{TRK}!$D$5:$D${last}'; RNG_QTY=f'{TRK}!$Z$5:$Z${last}'; RNG_C=f'{TRK}!$AA$5:$AA${last}'; RNG_R=f'{TRK}!$AB$5:$AB${last}'; RNG_B=f'{TRK}!$O$5:$O${last}'

# ---------- 금액 요약 ----------
ws=wb.create_sheet("금액 요약"); ws.sheet_view.showGridLines=False
h(ws,'B2',"27FW 구분별 금액 요약 — 원가합 · 소매가합",14)
ws['B3']="구분: 의류(J/T/V 등) · 용품(품번 Q로 시작) · WM/피싱(품번 8로 시작, 27FW 현재 0건). 트래커 값 수정 시 자동 갱신됩니다."
ws['B3'].font=Font(name=F,size=10,color="666666")
hdrs=["구분","스타일 수","기획수량 (장)","원가합 (원)","소매가합 (원)","소매가합 비중"]
for ci,t in enumerate(hdrs,start=2):
    c=ws.cell(row=5,column=ci,value=t)
    c.font=Font(name=F,size=10,bold=True,color="FFFFFF"); c.fill=PatternFill('solid',fgColor=BRAND)
    c.alignment=Alignment(horizontal='center'); c.border=border
Q='LEFT('+RNG_D+',1)="Q"'; E8='LEFT('+RNG_D+',1)="8"'
NOTQ8=f'(LEFT({RNG_D},1)<>"Q")*(LEFT({RNG_D},1)<>"8")*({RNG_D}<>"")'
cats=[("의류",NOTQ8),("용품 (Q)",f'({Q})*1'),("WM/피싱 (8)",f'({E8})*1')]
for ri,(name,cond) in enumerate(cats,start=6):
    ws[f'B{ri}']=name
    ws[f'C{ri}']=f'=SUMPRODUCT({cond})'
    ws[f'D{ri}']=f'=SUMPRODUCT({cond}*{RNG_QTY})'
    ws[f'E{ri}']=f'=SUMPRODUCT({cond}*{RNG_C})'
    ws[f'F{ri}']=f'=SUMPRODUCT({cond}*{RNG_R})'
    ws[f'G{ri}']=f'=IF($F$9=0,"",F{ri}/$F$9)'
tot=9
ws[f'B{tot}']="합계"
ws[f'C{tot}']='=SUM(C6:C8)'; ws[f'D{tot}']='=SUM(D6:D8)'; ws[f'E{tot}']='=SUM(E6:E8)'; ws[f'F{tot}']='=SUM(F6:F8)'; ws[f'G{tot}']='=IF(F9=0,"",1)'
for ri in range(6,10):
    for col in 'BCDEFG':
        c=ws[f'{col}{ri}']; c.font=Font(name=F,size=11,bold=(ri==tot)); c.border=border
        if col in 'CDEF': c.number_format='#,##0'
        if col=='G': c.number_format='0.0%'
        if col!='B': c.alignment=Alignment(horizontal='right')
    if ri==tot:
        for col in 'BCDEFG':
            ws[f'{col}{ri}'].fill=PatternFill('solid',fgColor=TOTALG)
            ws[f'{col}{ri}'].font=Font(name=F,size=11,bold=True,color=BRAND)
h(ws,'B12',"의류 내 비수기 선정 비중 (27FW 비수기 수령 시 자동 계산)",13)
ws['B13']="비수기 = KS_비수기 선정 27FW 탭 N열 'O' 기준 — 트래커 '비수기 선정' 열 입력 시 갱신"
ws['B13'].font=Font(name=F,size=10,color="666666")
hdrs2=["구분","스타일 수","기획수량 (장)","원가합 (원)","소매가합 (원)"]
for ci,t in enumerate(hdrs2,start=2):
    c=ws.cell(row=15,column=ci,value=t)
    c.font=Font(name=F,size=10,bold=True,color="FFFFFF"); c.fill=PatternFill('solid',fgColor=BRAND)
    c.alignment=Alignment(horizontal='center'); c.border=border
BIS=f'({RNG_B}="O")'
ws['B16']="의류 전체"; ws['C16']='=C6'; ws['D16']='=D6'; ws['E16']='=E6'; ws['F16']='=F6'
ws['B17']="비수기 선정 (의류)"
ws['C17']=f'=SUMPRODUCT({NOTQ8}*{BIS})'
ws['D17']=f'=SUMPRODUCT({NOTQ8}*{BIS}*{RNG_QTY})'
ws['E17']=f'=SUMPRODUCT({NOTQ8}*{BIS}*{RNG_C})'
ws['F17']=f'=SUMPRODUCT({NOTQ8}*{BIS}*{RNG_R})'
ws['B18']="비수기 비중"
for col in 'CDEF':
    ws[f'{col}18']=f'=IF({col}16=0,"",{col}17/{col}16)'
for ri in range(16,19):
    for col in 'BCDEF':
        c=ws[f'{col}{ri}']; c.font=Font(name=F,size=11); c.border=border
        if ri<18 and col in 'CDEF': c.number_format='#,##0'
        if ri==18 and col in 'CDEF': c.number_format='0.0%'
        if col!='B': c.alignment=Alignment(horizontal='right')
for col in 'BCDEF':
    ws[f'{col}18'].fill=PatternFill('solid',fgColor="DCEDF7"); ws[f'{col}18'].font=Font(name=F,size=11,bold=True,color=NAVY)
ws['B20']="* 원가합 = Σ(기획수량×입고원가), 소매가합 = Σ(기획수량×예상소매가) — 생진테 추출본 기준. 기획수량 미입력 스타일은 합계 미포함"
ws['B20'].font=Font(name=F,size=9,color="666666")
for col,wd in [('B',20),('C',11),('D',14),('E',18),('F',19),('G',13)]: ws.column_dimensions[col].width=wd

# ---------- 리스크 요약 ----------
ws=wb.create_sheet("리스크 요약"); ws.sheet_view.showGridLines=False
h(ws,'B2',"27FW 소재일정 리스크 요약 (생진테 출고일 · 띤쨘/르바란 보정)",14)
ws['B3']=(f'="기준일: "&TEXT(기준정보!$C$4,"YYYY-MM-DD")&"  ·  소재마감 = 출고일 − "&(기준정보!$C$5+기준정보!$C$6)&"일 − 휴무보정"'
          f'&"  ·  A "&COUNTIF({TRK}!$G$5:$G${last},"A(가을)")&" / W "&COUNTIF({TRK}!$G$5:$G${last},"W(겨울)")&" / X "&COUNTIF({TRK}!$G$5:$G${last},"X(사계절)")')
ws['B3'].font=Font(name=F,size=10,color="666666")
kpis=[("전체 스타일",f'=COUNTA({TRK}!$D$5:$D${last})',BRAND),
      ("A코드(가을) 스타일",f'=COUNTIF({TRK}!$G$5:$G${last},"A(가을)")',BRAND),
      ("★ 납기위험(A코드)",f'=COUNTIF({TRK}!$AE$5:$AE${last},"★*")',RED),
      ("◆ 지연주의(W/X)",f'=COUNTIF({TRK}!$AE$5:$AE${last},"◆*")',ORANGE),
      ("출고일 미정",f'=COUNTIF({TRK}!$AD$5:$AD${last},"출고일 미정")',NAVY)]
for i,(lab,f_,col_) in enumerate(kpis):
    col=get_column_letter(2+i*2)
    ws[f'{col}5']=lab; ws[f'{col}5'].font=Font(name=F,size=10,color="666666")
    ws[f'{col}6']=f_; ws[f'{col}6'].font=Font(name=F,size=22,bold=True,color=col_)
h(ws,'B9',"MD별 현황",12)
hdrs=["MD","스타일 수","A코드","출고일 미정","작지 미발행(X/부분)","비수기 선정","★ 납기위험","조율완료"]
for ci,t in enumerate(hdrs,start=2):
    c=ws.cell(row=10,column=ci,value=t)
    c.font=Font(name=F,size=10,bold=True,color="FFFFFF"); c.fill=PatternFill('solid',fgColor=BRAND)
    c.alignment=Alignment(horizontal='center'); c.border=border
mds=sorted({r['MD'] for r in rows if r['MD']})
for ri,md in enumerate(mds,start=11):
    ws[f'B{ri}']=md
    ws[f'C{ri}']=f'=COUNTIF({TRK}!$B$5:$B${last},$B{ri})'
    ws[f'D{ri}']=f'=COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$G$5:$G${last},"A(가을)")'
    ws[f'E{ri}']=f'=COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$AD$5:$AD${last},"출고일 미정")'
    ws[f'F{ri}']=f'=COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$P$5:$P${last},"X")+COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$P$5:$P${last},"부분")'
    ws[f'G{ri}']=f'=COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{RNG_B},"O")'
    ws[f'H{ri}']=f'=COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$AE$5:$AE${last},"★*")'
    ws[f'I{ri}']=f'=COUNTIFS({TRK}!$B$5:$B${last},$B{ri},{TRK}!$AF$5:$AF${last},"조율완료")'
    for col in 'BCDEFGHI':
        ws[f'{col}{ri}'].font=Font(name=F,size=10); ws[f'{col}{ri}'].border=border
        if col!='B': ws[f'{col}{ri}'].alignment=Alignment(horizontal='center')
tot2=11+len(mds)
ws[f'B{tot2}']="합계"; ws[f'B{tot2}'].font=Font(name=F,size=10,bold=True,color=BRAND)
for ci in range(3,10):
    cl=get_column_letter(ci)
    ws[f'{cl}{tot2}']=f'=SUM({cl}11:{cl}{tot2-1})'
    ws[f'{cl}{tot2}'].font=Font(name=F,size=10,bold=True,color=BRAND); ws[f'{cl}{tot2}'].alignment=Alignment(horizontal='center')
for col in 'BCDEFGHI':
    ws[f'{col}{tot2}'].border=border; ws[f'{col}{tot2}'].fill=PatternFill('solid',fgColor=TOTALG)
msg_r=tot2+3
h(ws,f'B{msg_r}',"공유 메시지 (관계부서 발송용)",12)
msg=("27FW 출고계획은 생진테 출고예정주차 기준으로 운영되며, 생산도 동 일정에 맞춰야 합니다. "
     "관건은 소재일정의 정확성입니다. 출고일 2주 전에는 제품이 입고되어야 하고, 제품입고 기준 100일 전에는 "
     "소재가 입고되어야 합니다(FW 봉제 난이도 반영). 미얀마 공장은 띤쨘 휴무(4/9~4/18), 인도네시아 공장은 "
     "르바란 휴무(3/4~3/17)로 생산기간이 겹치는 만큼 소재 마감이 추가로 앞당겨집니다. 이 마감까지 소재 입고가 "
     "확인되지 않았거나, 현시점 작업지시서가 발행되지 않은 A코드(가을상품)는 납기 준수가 어렵습니다. "
     "CMT는 본사(소재팀), 완사입은 생산처에서 소재 입고 가능일을 회신해 주시기 바랍니다.")
ws.merge_cells(f'B{msg_r+1}:I{msg_r+6}')
c=ws[f'B{msg_r+1}']; c.value=msg
c.font=Font(name=F,size=11); c.alignment=Alignment(wrap_text=True,vertical='top'); c.fill=PatternFill('solid',fgColor=PAPER)
for col,wd in [('B',16),('C',12),('D',12),('E',13),('F',18),('G',12),('H',13),('I',12)]: ws.column_dimensions[col].width=wd

wb.save('output/27FW_소재일정_역산_트래커.xlsx')
print("saved. styles:", N, "last:", last)
